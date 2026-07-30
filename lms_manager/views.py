import csv
import datetime
import io
import json
import re
from decimal import Decimal, InvalidOperation
from django.shortcuts import render, redirect, get_object_or_404
from django import forms
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_POST
from django.db.models import Sum, Count, Q
from django.db import transaction, DatabaseError
from django.contrib import messages
from django.utils import timezone
from .models import ClassRoom, Subject, Teacher, Student, Enrollment, Payment, PaymentPeriod, PaymentBatch, TeacherSettlement
from .forms import ClassRoomForm, SubjectForm, TeacherForm, StudentForm, EnrollmentForm, PaymentForm, PaymentPeriodForm

# -------------------------------------------------------------
# DASHBOARD
# -------------------------------------------------------------
def add_months(d, months):
    import datetime
    month = d.month - 1 + months
    year = d.year + month // 12
    month = month % 12 + 1
    day = min(d.day, [31,
        29 if (year % 4 == 0 and (year % 100 != 0 or year % 400 == 0)) else 28,
        31, 30, 31, 30, 31, 31, 30, 31, 30, 31][month-1])
    return datetime.date(year, month, day)


def assign_teacher_to_classroom(teacher, classroom):
    """Enforce the rule that one classroom is taught by one teacher."""
    classroom.teachers.set([teacher])
    subject = teacher.subjects.order_by('name').first()
    if not subject:
        return

    # Students in a class follow its assigned teacher.  Remove legacy
    # registrations for another teacher before creating the class enrollment.
    Enrollment.objects.filter(student__classroom=classroom).exclude(teacher=teacher).delete()
    for student in classroom.students.all():
        Enrollment.objects.get_or_create(
            student=student,
            subject=subject,
            defaults={'teacher': teacher},
        )


def dashboard_view(request):
    # Overall Counts
    student_count = Student.objects.count()
    teacher_count = Teacher.objects.count()
    subject_count = Subject.objects.count()
    class_count = ClassRoom.objects.count()
    period_count = PaymentPeriod.objects.count()
    payment_count = Payment.objects.count()

    # Total Revenue
    total_revenue = Payment.objects.aggregate(total=Sum('amount'))['total'] or 0

    # Revenue by payment period
    revenue_by_period = list(Payment.objects.values('payment_period__name')
                            .annotate(total=Sum('amount'))
                            .order_by('-total'))

    # Revenue by classroom
    revenue_by_class = list(Payment.objects.values('classroom__id', 'classroom__name')
                           .annotate(total=Sum('amount'))
                           .order_by('-total'))

    # Revenue by subject
    revenue_by_subject = list(Payment.objects.values('subject__name')
                             .annotate(total=Sum('amount'))
                             .order_by('-total'))

    # Revenue by teacher
    revenue_by_teacher = list(Payment.objects.values('teacher__id', 'teacher__name')
                             .annotate(total=Sum('amount'))
                             .order_by('-total'))

    # ---- Onboarding wizard state ----
    setup_steps = [
        {
            'key': 'subject',
            'icon': 'fa-book-open',
            'color': 'info',
            'label': 'Tạo Môn học',
            'desc': 'Thêm các môn (Toán, Vật lí, Hóa học...) để gán cho giảng viên.',
            'count': subject_count,
            'done': subject_count > 0,
            'url_name': 'subject_list',
        },
        {
            'key': 'teacher',
            'icon': 'fa-chalkboard-teacher',
            'color': 'warning',
            'label': 'Thêm Giảng viên',
            'desc': 'Tạo giảng viên và gán môn + lớp phụ trách.',
            'count': teacher_count,
            'done': teacher_count > 0,
            'url_name': 'teacher_list',
        },
    ]
    all_done = all(step['done'] for step in setup_steps)
    done_count = sum(1 for s in setup_steps if s['done'])

    # ---- Công nợ (nổi bật) ----
    # Tổng số bản ghi "chưa đóng" trong bảng kê
    debt_total_count = 0
    for period in PaymentPeriod.objects.all().select_related('subject', 'teacher'):
        unpaid = Enrollment.objects.filter(
            subject=period.subject,
            teacher=period.teacher,
        ).exclude(
            student__payments__payment_period=period
        )
        debt_total_count += unpaid.count()

    # Tính học sinh nợ từ 2 đợt trở lên
    today = timezone.localdate()
    overdue_students = []
    all_students = Student.objects.all().select_related('classroom').prefetch_related('enrollments', 'enrollments__subject', 'enrollments__teacher')
    
    for s in all_students:
        if not s.start_date:
            continue
        unpaid_overdue = []
        for i in range(1, 9):
            status = getattr(s, f"dot_{i}")
            if status == 'Chưa đóng':
                deadline = add_months(s.start_date, i)
                if today >= deadline:
                    unpaid_overdue.append(f"Đợt {i}")
        
        if len(unpaid_overdue) >= 2:
            enrollments = s.enrollments.all()
            subj_teacher_info = []
            for e in enrollments:
                subj_teacher_info.append(f"{e.subject.name} (GV: {e.teacher.name})")
            
            overdue_students.append({
                'student': s,
                'classroom': s.classroom,
                'overdue_count': len(unpaid_overdue),
                'periods_list': ", ".join(unpaid_overdue),
                'subjects': ", ".join(subj_teacher_info) if subj_teacher_info else "Chưa đăng ký môn"
            })
            
    # Sắp xếp theo số đợt nợ (nhiều trước), rồi tên học sinh
    overdue_students.sort(key=lambda x: (-x['overdue_count'], x['student'].name))

    context = {
        'student_count': student_count,
        'teacher_count': teacher_count,
        'subject_count': subject_count,
        'class_count': class_count,
        'period_count': period_count,
        'payment_count': payment_count,
        'total_revenue': total_revenue,
        'revenue_by_period': revenue_by_period,
        'revenue_by_class': revenue_by_class,
        'revenue_by_subject': revenue_by_subject,
        'revenue_by_teacher': revenue_by_teacher,
        'setup_steps': setup_steps,
        'setup_done_count': done_count,
        'setup_total': len(setup_steps),
        'setup_all_done': all_done,
        'debt_total_count': debt_total_count,
        'overdue_students': overdue_students,
    }
    return render(request, 'lms_manager/dashboard.html', context)


# -------------------------------------------------------------
# DEBT DASHBOARD (Công nợ)
# -------------------------------------------------------------
def debt_dashboard(request):
    """
    Tổng hợp tất cả các khoản học phí chưa đóng theo từng đợt/lớp/giảng viên.
    Có bộ lọc theo lớp, giảng viên, môn học.
    """
    classrooms = ClassRoom.objects.all().order_by('name')
    teachers = Teacher.objects.all().order_by('name')
    subjects = Subject.objects.all().order_by('name')

    classroom_id = request.GET.get('classroom_id', '').strip()
    teacher_id = request.GET.get('teacher_id', '').strip()
    subject_id = request.GET.get('subject_id', '').strip()

    periods = PaymentPeriod.objects.all().select_related('subject', 'teacher')

    if teacher_id:
        periods = periods.filter(teacher_id=teacher_id)
    if subject_id:
        periods = periods.filter(subject_id=subject_id)

    periods = periods.order_by('name', 'subject__name')

    period_reports = []
    total_debt_count = 0

    for period in periods:
        enrollment_filter = {
            'subject': period.subject,
            'teacher': period.teacher,
        }
        if classroom_id:
            enrollment_filter['student__classroom_id'] = classroom_id

        enrollments = Enrollment.objects.filter(
            **enrollment_filter
        ).exclude(
            student__payments__payment_period=period
        ).select_related('student', 'student__classroom').order_by('student__classroom__name', 'student__name')

        unpaid_records = []
        for e in enrollments:
            unpaid_records.append({
                'student': e.student,
                'classroom': e.student.classroom,
            })

        if not unpaid_records:
            continue

        total_debt_count += len(unpaid_records)

        period_reports.append({
            'period': period,
            'period_label': f"{period.name} - Môn: {period.subject.name} (GV: {period.teacher.name})",
            'records': unpaid_records,
        })

    context = {
        'classrooms': classrooms,
        'teachers': teachers,
        'subjects': subjects,
        'selected_classroom_id': classroom_id,
        'selected_teacher_id': teacher_id,
        'selected_subject_id': subject_id,
        'period_reports': period_reports,
        'total_debt_count': total_debt_count,
    }
    return render(request, 'lms_manager/debt_dashboard.html', context)


# -------------------------------------------------------------
# CLASSROOM CRUD
# -------------------------------------------------------------
def classroom_list(request):
    classes = ClassRoom.objects.all().prefetch_related('teachers')
    
    class_reports = []
    for c in classes:
        student_count = c.students.count()
        subject_count = Enrollment.objects.filter(student__classroom=c).values('subject').distinct().count()
        revenue = Payment.objects.filter(classroom=c).aggregate(total=Sum('amount'))['total'] or 0
        
        teachers_list = list(c.teachers.all())
        teachers_names = ", ".join(t.name for t in teachers_list)
        first_teacher_name = teachers_list[0].name if teachers_list else ""
        
        class_reports.append({
            'classroom': c,
            'student_count': student_count,
            'subject_count': subject_count,
            'revenue': revenue,
            'teachers_names': teachers_names,
            'first_teacher_name': first_teacher_name,
        })
        
    # Sort by classroom name, and then by first teacher name if names are identical
    class_reports.sort(key=lambda x: (x['classroom'].name.lower(), x['first_teacher_name'].lower()))
        
    return render(request, 'lms_manager/classroom_list.html', {'class_reports': class_reports})

def classroom_create(request):
    if request.method == 'POST':
        form = ClassRoomForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Đã thêm lớp học thành công!")
            return redirect('classroom_list')
    else:
        form = ClassRoomForm()
    return render(request, 'lms_manager/form.html', {'form': form, 'title': 'Thêm lớp học'})

def classroom_update(request, pk):
    classroom = get_object_or_404(ClassRoom, pk=pk)
    if request.method == 'POST':
        form = ClassRoomForm(request.POST, instance=classroom)
        if form.is_valid():
            form.save()
            messages.success(request, "Đã cập nhật lớp học thành công!")
            return redirect('classroom_detail', pk=classroom.id)
    else:
        form = ClassRoomForm(instance=classroom)
    return render(request, 'lms_manager/form.html', {'form': form, 'title': 'Sửa lớp học'})

def classroom_delete(request, pk):
    classroom = get_object_or_404(ClassRoom, pk=pk)
    classroom.delete()
    messages.success(request, "Đã xóa lớp học!")
    return redirect('classroom_list')


# -------------------------------------------------------------
# SUBJECT CRUD
# -------------------------------------------------------------
def subject_list(request):
    subjects = Subject.objects.all().order_by('name')
    return render(request, 'lms_manager/subject_list.html', {'subjects': subjects})

def subject_create(request):
    if request.method == 'POST':
        form = SubjectForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Đã thêm môn học thành công!")
            return redirect('subject_list')
    else:
        form = SubjectForm()
    return render(request, 'lms_manager/form.html', {'form': form, 'title': 'Thêm môn học'})

def subject_update(request, pk):
    subject = get_object_or_404(Subject, pk=pk)
    if request.method == 'POST':
        form = SubjectForm(request.POST, instance=subject)
        if form.is_valid():
            form.save()
            messages.success(request, "Đã cập nhật môn học thành công!")
            return redirect('subject_list')
    else:
        form = SubjectForm(instance=subject)
    return render(request, 'lms_manager/form.html', {'form': form, 'title': 'Sửa môn học'})

def subject_delete(request, pk):
    subject = get_object_or_404(Subject, pk=pk)
    subject.delete()
    messages.success(request, "Đã xóa môn học!")
    return redirect('subject_list')


# -------------------------------------------------------------
# TEACHER CRUD
# -------------------------------------------------------------
def teacher_list(request):
    search_query = request.GET.get('q', '').strip()
    teachers = Teacher.objects.all().prefetch_related('subjects', 'classes')
    if search_query:
        teachers = teachers.filter(
            Q(name__icontains=search_query) | Q(phone__icontains=search_query)
        )
    return render(request, 'lms_manager/teacher_list.html', {
        'teachers': teachers.order_by('name'),
        'search_query': search_query,
    })

def teacher_create(request):
    if request.method == 'POST':
        form = TeacherForm(request.POST)
        if form.is_valid():
            teacher = form.save(commit=False)
            
            # Capitalize teacher name (Title Case)
            teacher.name = " ".join(word.capitalize() for word in form.cleaned_data['name'].split())
            teacher.save()
            
            # Process typed subject names (supports comma-separated values)
            raw_subjects = form.cleaned_data.get('subject_name', '')
            subject_names = [s.strip() for s in raw_subjects.split(',') if s.strip()]
            
            subject_objs = []
            for name in subject_names:
                capitalized_sub = " ".join(word.capitalize() for word in name.split())
                sub, _ = Subject.objects.get_or_create(name=capitalized_sub)
                subject_objs.append(sub)
                
            teacher.subjects.set(subject_objs)
            
            # Process JSON-based classrooms list
            import json
            raw_classes_json = request.POST.get('classes_json', '[]')
            try:
                classes_data = json.loads(raw_classes_json)
            except Exception:
                classes_data = []
            
            class_objs = []
            for c_data in classes_data:
                class_id = c_data.get('id')
                class_name = c_data.get('name', '').strip()
                if not class_name:
                    continue
                
                capitalized_class = " ".join(
                    word.upper() if (len(word) <= 4 and any(char.isdigit() for char in word))
                    else word.capitalize()
                    for word in class_name.split()
                )
                
                if class_id:
                    try:
                        cls = ClassRoom.objects.get(pk=class_id)
                        if cls.name != capitalized_class:
                            cls.name = capitalized_class
                            cls.save()
                    except ClassRoom.DoesNotExist:
                        cls = ClassRoom.objects.create(name=capitalized_class)
                else:
                    cls = next((c for c in class_objs if c.name == capitalized_class), None)
                    if not cls:
                        cls = ClassRoom.objects.filter(name=capitalized_class).first()
                        if not cls:
                            cls = ClassRoom.objects.create(name=capitalized_class)
                
                class_objs.append(cls)
                
            previous_classes = list(teacher.classes.all())
            for classroom in previous_classes:
                if classroom not in class_objs:
                    classroom.teachers.remove(teacher)
            for classroom in class_objs:
                assign_teacher_to_classroom(teacher, classroom)
            
            messages.success(request, f"Đã thêm giảng viên '{teacher.name}' thành công!")
            return redirect('teacher_list')
    else:
        form = TeacherForm()
        
    existing_subjects = Subject.objects.all().order_by('name')
    existing_classrooms = ClassRoom.objects.all().order_by('name')
    return render(request, 'lms_manager/teacher_create_form.html', {
        'form': form,
        'title': 'Thêm giảng viên',
        'existing_subjects': existing_subjects,
        'existing_classrooms': existing_classrooms,
        'initial_classes_json': '[]'
    })


def teacher_update(request, pk):
    teacher = get_object_or_404(Teacher, pk=pk)
    if request.method == 'POST':
        form = TeacherForm(request.POST, instance=teacher)
        if form.is_valid():
            teacher = form.save(commit=False)
            
            # Capitalize teacher name (Title Case)
            teacher.name = " ".join(word.capitalize() for word in form.cleaned_data['name'].split())
            teacher.save()
            
            # Process typed subject names (supports comma-separated values)
            raw_subjects = form.cleaned_data.get('subject_name', '')
            subject_names = [s.strip() for s in raw_subjects.split(',') if s.strip()]
            
            subject_objs = []
            for name in subject_names:
                capitalized_sub = " ".join(word.capitalize() for word in name.split())
                sub, _ = Subject.objects.get_or_create(name=capitalized_sub)
                subject_objs.append(sub)
                
            teacher.subjects.set(subject_objs)
            
            # Process JSON-based classrooms list
            import json
            raw_classes_json = request.POST.get('classes_json', '[]')
            try:
                classes_data = json.loads(raw_classes_json)
            except Exception:
                classes_data = []
            
            class_objs = []
            for c_data in classes_data:
                class_id = c_data.get('id')
                class_name = c_data.get('name', '').strip()
                if not class_name:
                    continue
                
                capitalized_class = " ".join(
                    word.upper() if (len(word) <= 4 and any(char.isdigit() for char in word))
                    else word.capitalize()
                    for word in class_name.split()
                )
                
                if class_id:
                    try:
                        cls = ClassRoom.objects.get(pk=class_id)
                        if cls.name != capitalized_class:
                            cls.name = capitalized_class
                            cls.save()
                    except ClassRoom.DoesNotExist:
                        cls = ClassRoom.objects.create(name=capitalized_class)
                else:
                    cls = next((c for c in class_objs if c.name == capitalized_class), None)
                    if not cls:
                        cls = ClassRoom.objects.filter(name=capitalized_class).first()
                        if not cls:
                            cls = ClassRoom.objects.create(name=capitalized_class)
                
                class_objs.append(cls)
                
            previous_classes = list(teacher.classes.all())
            for classroom in previous_classes:
                if classroom not in class_objs:
                    classroom.teachers.remove(teacher)
            for classroom in class_objs:
                assign_teacher_to_classroom(teacher, classroom)
            
            messages.success(request, f"Đã cập nhật giảng viên '{teacher.name}' thành công!")
            return redirect('teacher_detail', pk=teacher.id)
    else:
        initial_data = {
            'subject_name': ", ".join(s.name for s in teacher.subjects.all()),
        }
        form = TeacherForm(instance=teacher, initial=initial_data)
        
    initial_classes = []
    for c in teacher.classes.all():
        initial_classes.append({'id': c.id, 'name': c.name})
    import json
    initial_classes_json = json.dumps(initial_classes)
        
    existing_subjects = Subject.objects.all().order_by('name')
    existing_classrooms = ClassRoom.objects.all().order_by('name')
    return render(request, 'lms_manager/teacher_create_form.html', {
        'form': form,
        'title': 'Sửa giảng viên',
        'existing_subjects': existing_subjects,
        'existing_classrooms': existing_classrooms,
        'initial_classes_json': initial_classes_json
    })


@require_POST
def teacher_delete(request, pk):
    teacher = get_object_or_404(Teacher, pk=pk)
    if TeacherSettlement.objects.filter(teacher=teacher).exists():
        messages.error(
            request,
            f"Không thể xóa giảng viên '{teacher.name}' vì đã có dữ liệu quyết toán. "
            "Hãy giữ giảng viên này để bảo toàn lịch sử tài chính.",
        )
        return redirect('teacher_list')

    teacher.delete()
    messages.success(request, "Đã xóa giảng viên!")
    return redirect('teacher_list')


@require_POST
def teacher_add_classroom(request, pk):
    teacher = get_object_or_404(Teacher, pk=pk)
    class_name = request.POST.get('class_name', '').strip()
    if class_name:
        capitalized_class = " ".join(
            word.upper() if (len(word) <= 4 and any(char.isdigit() for char in word))
            else word.capitalize()
            for word in class_name.split()
        )
        
        # Check duplicate for this specific teacher
        if teacher.classes.filter(name=capitalized_class).exists():
            messages.warning(request, f"Giảng viên '{teacher.name}' đã có lớp dạy tên '{capitalized_class}'!")
        else:
            cls = ClassRoom.objects.create(name=capitalized_class)
            assign_teacher_to_classroom(teacher, cls)
            messages.success(request, f"Đã thêm lớp dạy '{capitalized_class}' cho giảng viên '{teacher.name}'!")
    else:
        messages.error(request, "Tên lớp không hợp lệ!")
    return redirect('teacher_detail', pk=teacher.id)



def teacher_import_template(request):
    from django.http import HttpResponse
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="Template_Nhap_Lop_Giang_Vien.csv"'
    
    # Prepend UTF-8 BOM so Excel opens it with correct Vietnamese accents
    response.write(b'\xef\xbb\xbf')
    
    writer = csv.writer(response)
    writer.writerow(['Tên Lớp học', 'Môn giảng dạy'])
    writer.writerow(['10A1', 'Toán học'])
    writer.writerow(['11B2', 'Vật lí'])
    writer.writerow(['Lớp Toán nâng cao', 'Hình học'])
    return response


def teacher_import_excel(request, pk):
    teacher = get_object_or_404(Teacher, pk=pk)
    
    if request.method == 'POST':
        csv_file = request.FILES.get('excel_file')
        if not csv_file:
            messages.error(request, "Vui lòng chọn một tệp CSV để tải lên!")
            return redirect('teacher_import_excel', pk=teacher.id)
            
        if not csv_file.name.endswith('.csv'):
            messages.error(request, "Hệ thống chỉ hỗ trợ tệp định dạng .csv (mã hóa UTF-8)!")
            return redirect('teacher_import_excel', pk=teacher.id)
            
        try:
            # Read file content and decode from UTF-8 (utf-8-sig strips BOM)
            file_data = csv_file.read().decode('utf-8-sig')
            io_string = io.StringIO(file_data)
            
            # Read CSV
            reader = csv.reader(io_string, delimiter=',')
            header = next(reader, None) # skip header row
            
            imported_classes = 0
            imported_subjects = 0
            
            for row in reader:
                if not row or len(row) < 2:
                    continue
                
                raw_class = row[0].strip()
                raw_subject = row[1].strip()
                
                if not raw_class or not raw_subject:
                    continue
                
                # Format classroom name using our smart formatting
                capitalized_class = " ".join(
                    word.upper() if (len(word) <= 4 and any(char.isdigit() for char in word))
                    else word.capitalize()
                    for word in raw_class.split()
                )
                
                # Format subject name
                capitalized_sub = " ".join(word.capitalize() for word in raw_subject.split())
                
                # Get or create
                classroom = teacher.classes.filter(name=capitalized_class).first()
                class_created = False
                if not classroom:
                    classroom = ClassRoom.objects.create(name=capitalized_class)
                    teacher.classes.add(classroom)
                    class_created = True
                
                subject, sub_created = Subject.objects.get_or_create(name=capitalized_sub)
                
                if class_created:
                    imported_classes += 1
                if sub_created:
                    imported_subjects += 1
                    
                # Associate with teacher
                teacher.classes.add(classroom)
                teacher.subjects.add(subject)
                
            messages.success(
                request, 
                f"Nhập dữ liệu thành công! Đã liên kết giảng viên với các lớp & môn học tương ứng. "
                f"(Tạo mới {imported_classes} lớp học và {imported_subjects} môn học mới)."
            )
            return redirect('teacher_detail', pk=teacher.id)
            
        except Exception as e:
            messages.error(request, f"Đã xảy ra lỗi khi đọc tệp: {str(e)}")
            return redirect('teacher_import_excel', pk=teacher.id)
            
    return render(request, 'lms_manager/teacher_import_excel.html', {'teacher': teacher})


# -------------------------------------------------------------
# STUDENT & ENROLLMENT CRUD
# -------------------------------------------------------------
def student_list(request):
    classrooms = ClassRoom.objects.all().order_by('name')
    students = Student.objects.all().select_related('classroom').prefetch_related(
        'enrollments__subject', 'enrollments__teacher', 'payments__payment_period'
    )
    
    search_name = request.GET.get('name', '').strip()
    classroom_id = request.GET.get('classroom_id', '').strip()
    
    if search_name:
        students = students.filter(name__icontains=search_name)
    if classroom_id:
        students = students.filter(classroom_id=classroom_id)
        
    students = students.order_by('classroom__name', 'name')

    attach_payment_period_amounts(students)
    
    return render(request, 'lms_manager/student_list.html', {
        'students': students,
        'classrooms': classrooms,
        'search_name': search_name,
        'selected_classroom_id': classroom_id,
    })


def attach_payment_period_amounts(students, teacher=None):
    """Expose paid amounts by numbered period for read-only student tables."""
    period_pattern = re.compile(r'(?:Đợt|Period)\s+([1-8])', re.IGNORECASE)
    for student in students:
        amounts = {period_num: Decimal('0') for period_num in range(1, 9)}
        for payment in student.payments.all():
            if teacher and payment.teacher_id != teacher.id:
                continue
            match = period_pattern.fullmatch(payment.payment_period.name.strip())
            if match:
                period_num = int(match.group(1))
                amounts[period_num] += payment.amount
        for period_num, amount in amounts.items():
            setattr(student, f'payment_amount_{period_num}', amount or None)


def formatted_excel_response(title, headers, rows, filename, details=None):
    """Create a consistently formatted, single-page A4 XLSX download."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Danh sach'
    sheet.sheet_view.showGridLines = False
    column_count = len(headers)
    last_column = get_column_letter(column_count)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=column_count)
    title_cell = sheet['A1']
    title_cell.value = title
    title_cell.font = Font(bold=True, size=14, color='FFFFFF')
    title_cell.fill = PatternFill('solid', fgColor='1D4ED8')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    sheet.row_dimensions[1].height = 28

    next_row = 2
    for label, value in details or []:
        sheet.cell(row=next_row, column=1, value=label).font = Font(bold=True, color='334155')
        sheet.cell(row=next_row, column=2, value=value).font = Font(color='0F172A')
        next_row += 1
    if details:
        next_row += 1

    header_row = next_row
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='0F766E')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    sheet.row_dimensions[header_row].height = 28

    thin = Side(style='thin', color='D9E2F3')
    for row_index, row_values in enumerate(rows, start=header_row + 1):
        for column, value in enumerate(row_values, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value)
            cell.border = Border(bottom=thin)
            is_amount = isinstance(value, Decimal)
            cell.alignment = Alignment(
                horizontal='right' if is_amount else ('center' if column == 1 else 'left'),
                vertical='center',
                wrap_text=True,
            )
            if is_amount:
                cell.number_format = '#,##0'
        sheet.row_dimensions[row_index].height = 21

    for column in range(1, column_count + 1):
        values = [str(headers[column - 1])] + [
            '' if len(row) < column or row[column - 1] is None else str(row[column - 1])
            for row in rows
        ]
        width = min(max(max((len(value) for value in values), default=10) + 2, 10), 28)
        sheet.column_dimensions[get_column_letter(column)].width = width

    last_row = max(header_row + len(rows), header_row)
    sheet.auto_filter.ref = f'A{header_row}:{last_column}{last_row}'
    sheet.freeze_panes = f'A{header_row + 1}'
    sheet.print_area = f'A1:{last_column}{last_row}'
    sheet.print_title_rows = f'1:{header_row}'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = sheet.ORIENTATION_LANDSCAPE
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.4, bottom=0.4, header=0.15, footer=0.15)
    sheet.print_options.horizontalCentered = True

    buffer = io.BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f"attachment; filename*=UTF-8''{filename}"
    return response


def student_list_export(request):
    search_name = request.GET.get('name', '').strip()
    classroom_id = request.GET.get('classroom_id', '').strip()
    students = Student.objects.all().select_related('classroom').prefetch_related('payments__payment_period')
    if search_name:
        students = students.filter(name__icontains=search_name)
    if classroom_id:
        students = students.filter(classroom_id=classroom_id)
    students = list(students.order_by('classroom__name', 'name'))
    attach_payment_period_amounts(students)
    rows = []
    for index, student in enumerate(students, start=1):
        rows.append([
            index, student.name, student.phone or '', student.classroom.name if student.classroom else '',
            student.start_date.strftime('%d/%m/%Y') if student.start_date else '',
            *[getattr(student, f'payment_amount_{period}') or 'Chưa đóng' for period in range(1, 9)],
        ])
    return formatted_excel_response(
        'DANH SÁCH HỌC SINH',
        ['STT', 'Học sinh', 'SĐT', 'Lớp', 'Bắt đầu', *[f'Đợt {period}' for period in range(1, 9)]],
        rows,
        'danh_sach_hoc_sinh.xlsx',
    )


def classroom_list_export(request):
    rows = []
    for index, classroom in enumerate(ClassRoom.objects.all().prefetch_related('teachers').order_by('name'), start=1):
        rows.append([
            index,
            classroom.name,
            ', '.join(teacher.name for teacher in classroom.teachers.all()),
            classroom.students.count(),
            Enrollment.objects.filter(student__classroom=classroom).values('subject').distinct().count(),
            Payment.objects.filter(classroom=classroom).aggregate(total=Sum('amount'))['total'] or 0,
        ])
    return formatted_excel_response(
        'DANH SÁCH LỚP HỌC',
        ['STT', 'Lớp học', 'Giảng viên', 'Sĩ số', 'Số môn', 'Tổng học phí đã thu'],
        rows,
        'danh_sach_lop_hoc.xlsx',
    )


def classroom_export(request, pk):
    classroom = get_object_or_404(ClassRoom, pk=pk)
    teacher_id = request.GET.get('teacher')
    teacher = get_object_or_404(Teacher, pk=teacher_id, classes=classroom) if teacher_id else None
    students = Student.objects.filter(classroom=classroom).prefetch_related('payments__payment_period')
    if teacher:
        student_ids = Enrollment.objects.filter(
            student__classroom=classroom, teacher=teacher
        ).values_list('student_id', flat=True)
        students = students.filter(id__in=student_ids)
    students = list(students.order_by('name'))
    attach_payment_period_amounts(students, teacher=teacher)
    subject_by_student = {}
    enrollments = Enrollment.objects.filter(student__in=students)
    if teacher:
        enrollments = enrollments.filter(teacher=teacher)
    for enrollment in enrollments.select_related('subject'):
        subject_by_student.setdefault(enrollment.student_id, []).append(enrollment.subject.name)
    rows = []
    for index, student in enumerate(students, start=1):
        rows.append([
            index, student.name, student.phone or '', ', '.join(subject_by_student.get(student.id, [])),
            student.start_date.strftime('%d/%m/%Y') if student.start_date else '',
            *[getattr(student, f'payment_amount_{period}') or 'Chưa đóng' for period in range(1, 9)],
        ])
    details = [('Lớp', classroom.name)]
    if teacher:
        details.append(('Giảng viên', teacher.name))
    return formatted_excel_response(
        f'DANH SÁCH HỌC SINH - {classroom.name}',
        ['STT', 'Học sinh', 'SĐT', 'Môn học', 'Bắt đầu', *[f'Đợt {period}' for period in range(1, 9)]],
        rows,
        f'danh_sach_lop_{classroom.id}.xlsx',
        details,
    )


def student_export(request, pk):
    student = get_object_or_404(Student, pk=pk)
    payments = Payment.objects.filter(student=student).select_related('payment_period', 'subject', 'teacher').order_by('payment_period__name')
    rows = [
        [
            index, payment.payment_period.name, payment.subject.name, payment.teacher.name,
            payment.payment_date.strftime('%d/%m/%Y'), payment.amount,
        ]
        for index, payment in enumerate(payments, start=1)
    ]
    return formatted_excel_response(
        f'BẢNG KÊ HỌC PHÍ - {student.name}',
        ['STT', 'Đợt đóng tiền', 'Môn học', 'Giảng viên', 'Ngày đóng', 'Số tiền'],
        rows,
        f'bang_ke_hoc_phi_{student.id}.xlsx',
        [('Học sinh', student.name), ('Lớp', student.classroom.name if student.classroom else '')],
    )

def student_create(request):
    next_url = request.GET.get('next') or request.POST.get('next', '')
    
    if request.method == 'POST':
        form = StudentForm(request.POST)
        if form.is_valid():
            student = form.save()
            classroom_teacher = student.classroom.teachers.order_by('name').first() if student.classroom else None
            teacher = classroom_teacher or form.cleaned_data.get('teacher')
            subject = teacher.subjects.order_by('name').first() if teacher else form.cleaned_data.get('subject')

            if teacher and student.classroom and not classroom_teacher:
                assign_teacher_to_classroom(teacher, student.classroom)
            
            if subject and teacher:
                Enrollment.objects.get_or_create(
                    student=student,
                    subject=subject,
                    defaults={'teacher': teacher}
                )
                messages.success(request, f"Đã thêm học sinh '{student.name}' và đăng ký học môn {subject.name} với GV {teacher.name}!")
            else:
                messages.success(request, f"Đã thêm học sinh '{student.name}' thành công!")
                
            if next_url:
                return redirect(next_url)
            return redirect('student_list')
    else:
        initial = {}
        classroom_id = request.GET.get('classroom')
        if classroom_id:
            initial['classroom'] = classroom_id
            classroom = ClassRoom.objects.filter(id=classroom_id).first()
            if classroom:
                assigned_teacher = classroom.teachers.order_by('name').first()
                if assigned_teacher:
                    initial['teacher'] = assigned_teacher.id
                    subject = assigned_teacher.subjects.order_by('name').first()
                    if subject:
                        initial['subject'] = subject.id
            
        teacher_id = request.GET.get('teacher')
        if teacher_id:
            teacher = Teacher.objects.filter(id=teacher_id).first()
            if teacher:
                initial['teacher'] = teacher.id
                subject = teacher.subjects.first()
                if subject:
                    initial['subject'] = subject.id
                    
        form = StudentForm(initial=initial)
        
    return render(request, 'lms_manager/student_create_form.html', {
        'form': form,
        'next': next_url
    })

def student_update(request, pk):
    student = get_object_or_404(Student, pk=pk)
    if request.method == 'POST':
        form = StudentForm(request.POST, instance=student)
        if form.is_valid():
            form.save()
            messages.success(request, "Đã cập nhật học sinh thành công!")
            return redirect('student_detail', pk=student.id)
    else:
        form = StudentForm(instance=student)
    return render(request, 'lms_manager/form.html', {'form': form, 'title': 'Sửa học sinh'})

def student_delete(request, pk):
    student = get_object_or_404(Student, pk=pk)
    student.delete()
    messages.success(request, "Đã xóa học sinh!")
    return redirect('student_list')

def student_register(request, student_id):
    student = get_object_or_404(Student, pk=student_id)
    enrollments = Enrollment.objects.filter(student=student)
    
    if request.method == 'POST':
        form = EnrollmentForm(request.POST)
        if form.is_valid():
            # Check unique_together manually to provide friendly warning if already registered
            subject = form.cleaned_data['subject']
            if Enrollment.objects.filter(student=student, subject=subject).exists():
                messages.warning(request, f"Học sinh đã đăng ký môn {subject.name} trước đó rồi!")
            else:
                enrollment = form.save(commit=False)
                enrollment.student = student
                enrollment.save()
                
                # Update student's classroom
                classroom = form.cleaned_data['classroom']
                student.classroom = classroom
                student.save()
                
                messages.success(request, f"Đăng ký thành công môn {subject.name} cho {student.name}!")
            return redirect('student_register', student_id=student.id)
    else:
        form = EnrollmentForm(initial={'student': student, 'classroom': student.classroom})
        # Hide the student selection field or lock it since we are on the student-specific page
        form.fields['student'].widget = forms.HiddenInput()
        
    return render(request, 'lms_manager/student_register.html', {
        'student': student,
        'enrollments': enrollments,
        'form': form
    })

def enrollment_delete(request, pk):
    enrollment = get_object_or_404(Enrollment, pk=pk)
    student_id = enrollment.student.id
    enrollment.delete()
    messages.success(request, "Đã hủy đăng ký môn học!")
    return redirect('student_register', student_id=student_id)


# -------------------------------------------------------------
# PAYMENT & BILLING CRUD
# -------------------------------------------------------------
def payment_list(request):
    payments = Payment.objects.all().order_by('-payment_date', '-id')
    return render(request, 'lms_manager/payment_list.html', {'payments': payments})


def _payment_period_number(period_name):
    match = re.fullmatch(r'Đợt\s+([1-8])', period_name.strip())
    return int(match.group(1)) if match else None


def _create_batch_payments(items, payment_date, payment_method='cash'):
    """Create individual payments and their one combined customer receipt."""
    if not isinstance(items, list) or not items:
        raise ValueError('Hãy thêm ít nhất một khoản thu.')
    if payment_method not in dict(Payment.PAYMENT_METHOD_CHOICES):
        raise ValueError('Hình thức thanh toán không hợp lệ.')

    seen_items = set()
    prepared_items = []
    for index, item in enumerate(items, start=1):
        try:
            student = Student.objects.get(pk=int(item['student_id']))
            teacher = Teacher.objects.get(pk=int(item['teacher_id']))
            amount = Decimal(str(item['amount']))
        except (KeyError, TypeError, ValueError, InvalidOperation, Student.DoesNotExist, Teacher.DoesNotExist):
            raise ValueError(f'Khoản thu dòng {index} không hợp lệ.')

        raw_periods = item.get('period_nums', [item.get('period_num')])
        if not isinstance(raw_periods, list):
            raw_periods = [raw_periods]
        try:
            period_nums = [int(period_num) for period_num in raw_periods]
        except (TypeError, ValueError):
            raise ValueError(f'Khoản thu dòng {index} không hợp lệ.')

        if not period_nums or any(period_num not in range(1, 9) for period_num in period_nums) or amount <= 0:
            raise ValueError(f'Khoản thu dòng {index} không hợp lệ.')
        if not student.classroom:
            raise ValueError(f'Học sinh {student.name} chưa được xếp lớp.')

        enrollment = Enrollment.objects.filter(student=student, teacher=teacher).select_related('subject').first()
        if not enrollment:
            raise ValueError(f'Học sinh {student.name} chưa đăng ký học với giảng viên {teacher.name}.')

        for period_num in period_nums:
            item_key = (student.id, teacher.id, period_num)
            if item_key in seen_items:
                raise ValueError(f'Khoản thu dòng {index} bị trùng học sinh, giảng viên và đợt thu.')
            seen_items.add(item_key)
            if Payment.objects.filter(
                student=student,
                teacher=teacher,
                payment_period__name=f'Đợt {period_num}',
            ).exists():
                raise ValueError(
                    f'Học sinh {student.name} đã đóng Đợt {period_num} với giảng viên {teacher.name}.'
                )
            prepared_items.append((student, teacher, enrollment.subject, period_num, amount))

    with transaction.atomic():
        batch = PaymentBatch.objects.create(payment_date=payment_date)
        payments = []
        for student, teacher, subject, period_num, amount in prepared_items:
            period, _ = PaymentPeriod.objects.get_or_create(
                name=f'Đợt {period_num}', teacher=teacher, subject=subject
            )
            payment = Payment.objects.create(
                student=student,
                classroom=student.classroom,
                subject=subject,
                teacher=teacher,
                payment_period=period,
                amount=amount,
                payment_date=payment_date,
                payment_method=payment_method,
            )
            # Retain the legacy status for pages that still display it.  The
            # student list itself derives its status from Payment records.
            setattr(student, f'dot_{period_num}', 'Đã đóng')
            student.save(update_fields=[f'dot_{period_num}'])
            payments.append(payment)

        batch.payments.add(*payments)
        batch.generate_receipt_pdf()
    return batch

def payment_create(request):
    if request.method == 'POST':
        batch_items = request.POST.get('batch_items')
        if batch_items is not None:
            try:
                items = json.loads(batch_items)
                payment_date = datetime.datetime.strptime(
                    request.POST.get('payment_date', ''), '%Y-%m-%d'
                ).date()
                batch = _create_batch_payments(
                    items,
                    payment_date,
                    request.POST.get('payment_method', 'cash'),
                )
            except DatabaseError:
                error = 'Cơ sở dữ liệu chưa được cập nhật cho tính năng thu học phí. Hãy chạy lệnh migrate rồi thử lại.'
                if request.headers.get('x-requested-with') == 'XMLHttpRequest':
                    return JsonResponse({'success': False, 'error': error}, status=503)
                messages.error(request, error)
            except (ValueError, json.JSONDecodeError) as exc:
                error = str(exc) or 'Dữ liệu thu học phí không hợp lệ.'
                if request.headers.get('x-requested-with') == 'XMLHttpRequest':
                    return JsonResponse({'success': False, 'error': error}, status=400)
                messages.error(request, error)
            else:
                from django.urls import reverse
                receipt_url = reverse('payment_batch_receipt', kwargs={'pk': batch.id})
                if request.headers.get('x-requested-with') == 'XMLHttpRequest':
                    return JsonResponse({
                        'success': True,
                        'receipt_url': receipt_url,
                        'redirect_url': reverse('payment_list'),
                    })
                return redirect(receipt_url)

        form = PaymentForm(request.POST)
        if form.is_valid():
            payment = form.save(commit=False)
            
            # Get typed student name and parse name, classroom, phone
            raw_name = form.cleaned_data.get('student_name', '').strip()
            
            match = re.match(r"^(.+?)\s*\(\s*Lớp\s+(.+?)(?:\s*-\s*SĐT:\s*(.+?))?\s*\)$", raw_name, re.IGNORECASE)
            if match:
                clean_name = match.group(1).strip()
                classroom_name = match.group(2).strip()
                phone = match.group(3).strip() if match.group(3) else None
            else:
                clean_name = raw_name
                classroom_name = None
                phone = None
                
            capitalized_name = " ".join(word.capitalize() for word in clean_name.split())
            
            student = None
            if classroom_name:
                student = Student.objects.filter(
                    name=capitalized_name,
                    classroom__name__iexact=classroom_name
                )
                if phone:
                    student = student.filter(phone=phone)
                student = student.first()
            
            if not student:
                student = Student.objects.filter(name=capitalized_name).first()
                
            teacher = form.cleaned_data.get('teacher')
            
            created = False
            if not student:
                if classroom_name:
                    # Smart formatting for classroom
                    capitalized_class = " ".join(
                        word.upper() if (len(word) <= 4 and any(char.isdigit() for char in word))
                        else word.capitalize()
                        for word in classroom_name.split()
                    )
                    if teacher:
                        classroom = teacher.classes.filter(name=capitalized_class).first()
                    else:
                        classroom = ClassRoom.objects.filter(name=capitalized_class).first()
                    
                    if not classroom:
                        classroom = ClassRoom.objects.create(name=capitalized_class)
                        if teacher:
                            teacher.classes.add(classroom)
                else:
                    classroom = ClassRoom.objects.first() or ClassRoom.objects.create(name="Lớp Mới")
                
                student = Student.objects.create(
                    name=capitalized_name,
                    phone=phone,
                    classroom=classroom
                )
                created = True
            else:
                classroom = student.classroom
            
            # Resolve subject from teacher
            enrollment = Enrollment.objects.filter(student=student, teacher=teacher).first()
            if enrollment:
                subject = enrollment.subject
            else:
                subject = teacher.subjects.first() or Subject.objects.first()
                # Auto-enroll student in the subject if not enrolled yet
                Enrollment.objects.get_or_create(
                    student=student,
                    subject=subject,
                    defaults={'teacher': teacher}
                )
            
            # Parse the period name or index
            period_val = form.cleaned_data.get('payment_period', '').strip()
            if period_val.isdigit():
                period_num = int(period_val)
                period_name = f"Đợt {period_val}"
            else:
                period_name = period_val
                # Try to extract number
                period_num = None
                for i in range(1, 9):
                    if str(i) in period_name:
                        period_num = i
                        break
            
            # Find or create PaymentPeriod
            period_obj, _ = PaymentPeriod.objects.get_or_create(
                name=period_name,
                teacher=teacher,
                subject=subject
            )
            
            payment.student = student
            payment.classroom = classroom
            payment.subject = subject
            payment.teacher = teacher
            payment.payment_period = period_obj
            payment.save()
            
            # Update student status to "Đã đóng"
            if period_num:
                setattr(student, f"dot_{period_num}", "Đã đóng")
                student.save()
            
            if created:
                messages.success(request, f"Đã thêm học sinh mới '{student.name}' vào lớp {classroom.name}!")
            messages.success(request, f"Đã ghi nhận đóng tiền thành công đợt '{payment.payment_period.name}' cho học sinh '{student.name}'!")
            
            is_ajax = request.headers.get('x-requested-with') == 'XMLHttpRequest' or 'application/json' in request.META.get('HTTP_ACCEPT', '')
            if is_ajax:
                from django.urls import reverse
                return JsonResponse({
                    'success': True,
                    'receipt_url': reverse('payment_receipt', kwargs={'pk': payment.id}),
                    'redirect_url': reverse('payment_list')
                })
            return redirect('payment_receipt', pk=payment.id)
        else:
            is_ajax = request.headers.get('x-requested-with') == 'XMLHttpRequest' or 'application/json' in request.META.get('HTTP_ACCEPT', '')
            if is_ajax:
                return JsonResponse({
                    'success': False,
                    'errors': form.errors.get_json_data()
                })
            messages.error(request, "Thông tin nhập không hợp lệ!")
    else:
        initial_data = {}
        student_id = request.GET.get('student')
        if student_id:
            student = Student.objects.filter(id=student_id).first()
            if student:
                initial_data['student_name'] = student.name
        form = PaymentForm(initial=initial_data)
        
    existing_students = Student.objects.all().order_by('name').prefetch_related(
        'enrollments__teacher', 'payments__payment_period'
    )
    for student in existing_students:
        paid_periods = {
            int(match.group(1))
            for payment in student.payments.all()
            if (match := re.fullmatch(r'Đợt\s+([1-8])', payment.payment_period.name.strip()))
        }
        student.unpaid_period_numbers = [
            period_num for period_num in range(1, 9) if period_num not in paid_periods
        ]
    return render(request, 'lms_manager/payment_create.html', {
        'form': form,
        'existing_students': existing_students
    })


def payment_delete(request, pk):
    payment = get_object_or_404(Payment, pk=pk)
    student = payment.student
    period = payment.payment_period
    
    # Try to extract the period number
    period_num = None
    if period:
        for i in range(1, 9):
            if str(i) in period.name:
                period_num = i
                break
                
    payment.delete()
    
    if student and period_num:
        other_exists = Payment.objects.filter(student=student, payment_period=period).exists()
        if not other_exists:
            setattr(student, f"dot_{period_num}", "Chưa đóng")
            student.save()
            
    messages.success(request, "Đã xóa đợt đóng tiền!")
    
    next_url = request.GET.get('next') or request.META.get('HTTP_REFERER')
    if next_url:
        return redirect(next_url)
    return redirect('payment_list')


def payment_edit(request, pk):
    payment = get_object_or_404(Payment, pk=pk)
    old_student = payment.student
    old_period = payment.payment_period
    
    if request.method == 'POST':
        form = PaymentForm(request.POST, instance=payment)
        if form.is_valid():
            raw_name = form.cleaned_data.get('student_name', '').strip()
            match = re.match(r"^(.+?)\s*\(\s*Lớp\s+(.+?)(?:\s*-\s*SĐT:\s*(.+?))?\s*\)$", raw_name, re.IGNORECASE)
            if match:
                clean_name = match.group(1).strip()
                classroom_name = match.group(2).strip()
                phone = match.group(3).strip() if match.group(3) else None
            else:
                clean_name = raw_name
                classroom_name = None
                phone = None
                
            capitalized_name = " ".join(word.capitalize() for word in clean_name.split())
            
            student = None
            if classroom_name:
                student = Student.objects.filter(
                    name=capitalized_name,
                    classroom__name__iexact=classroom_name
                )
                if phone:
                    student = student.filter(phone=phone)
                student = student.first()
            if not student:
                student = Student.objects.filter(name=capitalized_name).first()
            if not student:
                student = old_student
                
            teacher = form.cleaned_data.get('teacher')
            
            enrollment = Enrollment.objects.filter(student=student, teacher=teacher).first()
            if enrollment:
                subject = enrollment.subject
            else:
                subject = teacher.subjects.first() or Subject.objects.first()
                Enrollment.objects.get_or_create(
                    student=student,
                    subject=subject,
                    defaults={'teacher': teacher}
                )
                
            period_val = form.cleaned_data.get('payment_period', '').strip()
            if period_val.isdigit():
                period_num = int(period_val)
                period_name = f"Đợt {period_val}"
            else:
                period_name = period_val
                period_num = None
                for i in range(1, 9):
                    if str(i) in period_name:
                        period_num = i
                        break
                        
            period_obj, _ = PaymentPeriod.objects.get_or_create(
                name=period_name,
                teacher=teacher,
                subject=subject
            )
            
            if old_period:
                old_num = None
                for i in range(1, 9):
                    if str(i) in old_period.name:
                        old_num = i
                        break
                if old_num:
                    other_exists = Payment.objects.filter(student=old_student, payment_period=old_period).exclude(id=payment.id).exists()
                    if not other_exists:
                        setattr(old_student, f"dot_{old_num}", "Chưa đóng")
                        old_student.save()
            
            if period_num:
                setattr(student, f"dot_{period_num}", "Đã đóng")
                student.save()
                
            payment.student = student
            payment.classroom = student.classroom
            payment.subject = subject
            payment.teacher = teacher
            payment.payment_period = period_obj
            payment.save()
            
            payment.generate_receipt_pdf()
            
            messages.success(request, "Đã cập nhật đợt đóng học phí thành công!")
            return redirect('student_detail', pk=student.id)
    else:
        initial_data = {}
        if payment.student:
            initial_data['student_name'] = f"{payment.student.name} (Lớp {payment.student.classroom.name})"
        if payment.payment_period:
            period_name = payment.payment_period.name
            # If name is "Đợt N", extract N
            if period_name.startswith("Đợt "):
                initial_data['payment_period'] = period_name.replace("Đợt ", "").strip()
            else:
                initial_data['payment_period'] = period_name
            
        form = PaymentForm(instance=payment, initial=initial_data)
        
    existing_students = Student.objects.all().order_by('name')
    return render(request, 'lms_manager/payment_edit_form.html', {
        'form': form,
        'payment': payment,
        'existing_students': existing_students
    })


from django.views.decorators.clickjacking import xframe_options_sameorigin

@xframe_options_sameorigin
def payment_receipt(request, pk):
    payment = get_object_or_404(Payment, pk=pk)
    import os
    if (
        not payment.receipt_pdf
        or not os.path.exists(payment.receipt_pdf.path)
        or not payment.receipt_pdf.name.endswith('_a5_v2.pdf')
    ):
        payment.generate_receipt_pdf()
        payment.refresh_from_db()
        
    response = HttpResponse(payment.receipt_pdf.read(), content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="receipt_{payment.id}.pdf"'
    return response


@xframe_options_sameorigin
def payment_batch_receipt(request, pk):
    batch = get_object_or_404(PaymentBatch, pk=pk)
    import os
    if (
        not batch.receipt_pdf
        or not os.path.exists(batch.receipt_pdf.path)
        or not batch.receipt_pdf.name.endswith('_a5_v2.pdf')
    ):
        batch.generate_receipt_pdf()
        batch.refresh_from_db()

    response = HttpResponse(batch.receipt_pdf.read(), content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="batch_receipt_{batch.id}.pdf"'
    return response


# -------------------------------------------------------------
# AJAX API FOR DYNAMIC DROPDOWNS
# -------------------------------------------------------------
def get_student_subjects(request, student_id):
    enrollments = Enrollment.objects.filter(student_id=student_id).select_related('subject')
    data = []
    for e in enrollments:
        data.append({
            'id': e.subject.id,
            'name': e.subject.name
        })
    return JsonResponse({'subjects': data})


def get_student_payment_periods(request, student_id):
    student = get_object_or_404(Student, pk=student_id)
    
    # Get subjects student is enrolled in
    enrollments = student.enrollments.all()
    
    # Construct filters for student's registered subject-teacher pairs
    from django.db.models import Q
    q_filter = Q()
    for enrollment in enrollments:
        q_filter |= Q(subject_id=enrollment.subject_id, teacher_id=enrollment.teacher_id)
        
    periods = []
    if enrollments:
        periods = list(
            PaymentPeriod.objects.filter(q_filter)
            .select_related('subject', 'teacher')
            .order_by('name', 'subject__name')
        )
    
    data = []
    for p in periods:
        data.append({
            'id': p.id,
            'name': f"{p.name} - Môn: {p.subject.name} (GV: {p.teacher.name})"
        })
        
    return JsonResponse({'periods': data})


def filter_payment_periods(request):
    student_id = request.GET.get('student_id')
    teacher_id = request.GET.get('teacher_id')
    
    periods = PaymentPeriod.objects.all().select_related('subject', 'teacher')
    
    if student_id:
        student = Student.objects.filter(id=student_id).first()
        if student:
            enrollments = student.enrollments.all()
            from django.db.models import Q
            q_filter = Q()
            for enrollment in enrollments:
                q_filter |= Q(subject_id=enrollment.subject_id, teacher_id=enrollment.teacher_id)
            if enrollments:
                periods = periods.filter(q_filter)
            else:
                periods = periods.none()
            
    if teacher_id:
        periods = periods.filter(teacher_id=teacher_id)
        
    periods = periods.order_by('name', 'subject__name')
    
    data = []
    for p in periods:
        data.append({
            'id': p.id,
            'name': f"{p.name} - Môn: {p.subject.name} (GV: {p.teacher.name})"
        })
        
    return JsonResponse({'periods': data})


def get_subject_teachers(request, subject_id):
    teachers = Teacher.objects.filter(subjects__id=subject_id).distinct().order_by('name')
    data = [{'id': t.id, 'name': t.name} for t in teachers]
    return JsonResponse({'teachers': data})


def get_teacher_classrooms(request, teacher_id):
    teacher = get_object_or_404(Teacher, pk=teacher_id)
    classrooms = teacher.classes.all().order_by('name')
    data = [{'id': c.id, 'name': c.name} for c in classrooms]
    return JsonResponse({'classrooms': data})


def get_teacher_subjects(request, teacher_id):
    teacher = get_object_or_404(Teacher, pk=teacher_id)
    subjects = teacher.subjects.all().order_by('name')
    data = [{'id': s.id, 'name': s.name} for s in subjects]
    return JsonResponse({'subjects': data})


def get_period_details(request, period_id):
    period = get_object_or_404(PaymentPeriod, pk=period_id)
    last_payment = Payment.objects.filter(subject=period.subject).order_by('-id').first()
    
    if last_payment:
        amount = int(last_payment.amount)
    else:
        sub_name = period.subject.name.lower()
        if "toán" in sub_name:
            amount = 500000
        elif "văn" in sub_name:
            amount = 450000
        elif "anh" in sub_name:
            amount = 600000
        elif "lý" in sub_name or "lí" in sub_name or "hóa" in sub_name:
            amount = 550000
        else:
            amount = 500000
            
    return JsonResponse({
        'amount': amount,
        'teacher_id': period.teacher.id,
    })


def teacher_detail(request, pk):
    teacher = get_object_or_404(Teacher, pk=pk)
    subjects = teacher.subjects.all()
    
    # Get classrooms associated with this teacher
    classrooms = teacher.classes.all().order_by('name')
    
    classroom_reports = []
    for cls in classrooms:
        # Find students in this classroom enrolled under this teacher
        enrollments = Enrollment.objects.filter(
            student__classroom=cls,
            teacher=teacher
        ).select_related('student', 'subject').prefetch_related(
            'student__payments__payment_period'
        ).order_by('student__name')
        
        students_list = []
        for e in enrollments:
            student = e.student
            student.enrolled_subject = e.subject.name
            students_list.append(student)

        attach_payment_period_amounts(students_list, teacher=teacher)
            
        classroom_reports.append({
            'classroom': cls,
            'students': students_list,
            'count': len(students_list)
        })
        
    total_earned = Payment.objects.filter(teacher=teacher).aggregate(total=Sum('amount'))['total'] or 0
    total_net_earned = float(total_earned) * 0.8
    total_settled = TeacherSettlement.objects.filter(teacher=teacher).aggregate(total=Sum('teacher_amount'))['total'] or 0
    total_pending_settlement = Decimal(str(total_net_earned)) - total_settled
    
    return render(request, 'lms_manager/teacher_detail.html', {
        'teacher': teacher,
        'subjects': subjects,
        'classroom_reports': classroom_reports,
        'total_earned': total_earned,
        'total_net_earned': total_net_earned,
        'total_settled': total_settled,
        'total_pending_settlement': max(total_pending_settlement, Decimal('0')),
        'existing_classrooms': ClassRoom.objects.all().order_by('name'),
    })


@require_POST
def teacher_classroom_student_remove(request, teacher_id, classroom_id, student_id):
    """Remove a student from a teacher's class without deleting their record or payments."""
    teacher = get_object_or_404(Teacher, pk=teacher_id)
    classroom = get_object_or_404(ClassRoom, pk=classroom_id, teachers=teacher)
    student = get_object_or_404(Student, pk=student_id, classroom=classroom)

    Enrollment.objects.filter(student=student, teacher=teacher).delete()
    student.classroom = None
    student.save(update_fields=['classroom'])
    messages.success(request, f'Đã đưa {student.name} ra khỏi lớp {classroom.name}. Hồ sơ và các phiếu thu đã được giữ lại.')
    return redirect('teacher_detail', pk=teacher.pk)


TEACHER_SETTLEMENT_RATE = Decimal('0.80')
SETTLEMENT_PERIOD_NUMBERS = tuple(range(1, 9))
SETTLEMENT_PERIOD_PATTERN = re.compile(r'(?:Đợt|Period)\s+([1-8])', re.IGNORECASE)


def _payment_period_number(payment_period):
    match = SETTLEMENT_PERIOD_PATTERN.fullmatch(payment_period.name.strip())
    return int(match.group(1)) if match else None


def _teacher_class_payments(teacher, classroom, period_numbers=None, unsettled_only=False):
    """Return tuition records for a teacher/class, optionally only unsettled ones."""
    payments = Payment.objects.filter(
        teacher=teacher,
        classroom=classroom,
    ).select_related('student', 'payment_period').order_by('student__name', 'payment_date', 'id')
    if unsettled_only:
        payments = payments.filter(teacher_settlements__isnull=True)
    if period_numbers:
        selected_numbers = set(period_numbers)
        matching_period_ids = [
            period.id
            for period in PaymentPeriod.objects.filter(teacher=teacher)
            if _payment_period_number(period) in selected_numbers
        ]
        payments = payments.filter(payment_period_id__in=matching_period_ids)
    return payments


def _unsettled_teacher_payments(teacher, classroom, period_numbers=None):
    return _teacher_class_payments(teacher, classroom, period_numbers, unsettled_only=True)


def _settlement_student_rows(payments, period_numbers=None):
    rows = {}
    for payment in payments:
        row = rows.setdefault(payment.student_id, {
            'student_name': payment.student.name,
            'phone': payment.student.phone or '',
            'period_names': [],
            'amounts_by_period_number': {},
            'settled_amounts_by_period_number': {},
            'unsettled_amounts_by_period_number': {},
            'amount': Decimal('0'),
        })
        row['amount'] += payment.amount
        period_number = _payment_period_number(payment.payment_period)
        if period_number:
            row['amounts_by_period_number'][period_number] = (
                row['amounts_by_period_number'].get(period_number, Decimal('0')) + payment.amount
            )
            if payment.teacher_settlements.exists():
                row['settled_amounts_by_period_number'][period_number] = (
                    row['settled_amounts_by_period_number'].get(period_number, Decimal('0')) + payment.amount
                )
            else:
                row['unsettled_amounts_by_period_number'][period_number] = (
                    row['unsettled_amounts_by_period_number'].get(period_number, Decimal('0')) + payment.amount
                )
        if payment.payment_period.name not in row['period_names']:
            row['period_names'].append(payment.payment_period.name)
    ordered_rows = list(sorted(rows.values(), key=lambda row: row['student_name'].lower()))
    if period_numbers:
        for row in ordered_rows:
            row['period_states'] = []
            for period_number in period_numbers:
                if amount := row['unsettled_amounts_by_period_number'].get(period_number):
                    row['period_states'].append({'status': 'unsettled', 'amount': amount})
                elif amount := row['settled_amounts_by_period_number'].get(period_number):
                    row['period_states'].append({'status': 'settled', 'amount': amount})
                else:
                    row['period_states'].append({'status': 'unpaid', 'amount': None})
    return ordered_rows


def teacher_class_settlement(request, teacher_id, classroom_id):
    teacher = get_object_or_404(Teacher, pk=teacher_id)
    classroom = get_object_or_404(ClassRoom, pk=classroom_id, teachers=teacher)
    selected_period_numbers = [
        int(value) for value in request.GET.getlist('period')
        if value.isdigit() and int(value) in SETTLEMENT_PERIOD_NUMBERS
    ]
    period_options = [
        {'number': period_number, 'name': f'Đợt {period_number}'}
        for period_number in SETTLEMENT_PERIOD_NUMBERS
    ]

    if request.method == 'POST':
        selected_period_numbers = [
            int(value) for value in request.POST.getlist('period')
            if value.isdigit() and int(value) in SETTLEMENT_PERIOD_NUMBERS
        ]
        if not selected_period_numbers:
            messages.error(request, 'Hãy chọn ít nhất một đợt cần quyết toán.')
        else:
            with transaction.atomic():
                payment_ids = list(_unsettled_teacher_payments(
                    teacher, classroom, selected_period_numbers
                ).values_list('id', flat=True))
                # Lock only Payment rows.  The unsettled check uses a LEFT JOIN
                # through the settlement M2M table, which PostgreSQL cannot lock.
                payments = list(Payment.objects.select_for_update(of=('self',)).filter(
                    id__in=payment_ids,
                    teacher=teacher,
                    classroom=classroom,
                    teacher_settlements__isnull=True,
                ).select_related('student', 'payment_period').order_by('student__name', 'id'))
                if not payments:
                    messages.error(request, 'Các đợt đã chọn không còn khoản thu nào cần quyết toán.')
                else:
                    revenue = sum((payment.amount for payment in payments), Decimal('0'))
                    teacher_amount = (revenue * TEACHER_SETTLEMENT_RATE).quantize(Decimal('1'))
                    settlement = TeacherSettlement.objects.create(
                        teacher=teacher,
                        classroom=classroom,
                        revenue=revenue,
                        teacher_share_rate=TEACHER_SETTLEMENT_RATE,
                        teacher_amount=teacher_amount,
                    )
                    settlement.payments.add(*payments)
                    messages.success(request, 'Đã quyết toán cho giảng viên. File Excel đang được tải xuống.')
                    return redirect('teacher_settlement_export', pk=settlement.pk)

    selected_payments = list(_unsettled_teacher_payments(
        teacher, classroom, selected_period_numbers
    )) if selected_period_numbers else []
    display_payments = list(_teacher_class_payments(
        teacher, classroom, selected_period_numbers
    ).prefetch_related('teacher_settlements')) if selected_period_numbers else []
    revenue = sum((payment.amount for payment in selected_payments), Decimal('0'))
    teacher_amount = (revenue * TEACHER_SETTLEMENT_RATE).quantize(Decimal('1'))
    return render(request, 'lms_manager/teacher_settlement.html', {
        'teacher': teacher,
        'classroom': classroom,
        'period_options': period_options,
        'selected_period_numbers': selected_period_numbers,
        'student_rows': _settlement_student_rows(display_payments, SETTLEMENT_PERIOD_NUMBERS),
        'payment_count': len(selected_payments),
        'revenue': revenue,
        'teacher_amount': teacher_amount,
        'teacher_share_rate': TEACHER_SETTLEMENT_RATE * 100,
    })


def teacher_settlement_export(request, pk):
    settlement = get_object_or_404(
        TeacherSettlement.objects.select_related('teacher', 'classroom'), pk=pk
    )
    # openpyxl is already a project dependency and is used for the existing Excel import flow.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.page import PageMargins

    rows = _settlement_student_rows(settlement.payments.select_related(
        'student', 'payment_period'
    ).order_by('student__name', 'payment_date', 'id'))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Quyet toan'
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'A6'
    sheet.merge_cells('A1:E1')
    sheet['A1'] = f'QUYẾT TOÁN GIẢNG VIÊN - {settlement.teacher.name}'
    sheet['A1'].font = Font(bold=True, size=14, color='FFFFFF')
    sheet['A1'].fill = PatternFill('solid', fgColor='1D4ED8')
    sheet['A1'].alignment = Alignment(horizontal='center')
    sheet.row_dimensions[1].height = 28
    sheet.append(['Lớp', settlement.classroom.name])
    sheet.append(['Ngày quyết toán', timezone.localtime(settlement.settled_at).strftime('%d/%m/%Y %H:%M')])
    sheet.append([])
    for row_number in (2, 3):
        sheet.cell(row=row_number, column=1).font = Font(bold=True, color='334155')
        sheet.cell(row=row_number, column=2).font = Font(color='0F172A')
    header_row = 5
    headers = ['STT', 'Học sinh', 'SĐT', 'Đợt đã đóng', 'Số tiền đã đóng (VNĐ)']
    sheet.append(headers)
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='0F766E')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    sheet.row_dimensions[header_row].height = 26
    thin_border = Side(style='thin', color='D9E2F3')
    table_border = Border(bottom=thin_border)
    for index, row in enumerate(rows, start=1):
        sheet.append([index, row['student_name'], row['phone'], ', '.join(row['period_names']), int(row['amount'])])
        row_number = header_row + index
        sheet.row_dimensions[row_number].height = 21
        for column in range(1, 6):
            cell = sheet.cell(row=row_number, column=column)
            cell.border = table_border
            cell.alignment = Alignment(
                horizontal='right' if column == 5 else ('center' if column == 1 else 'left'),
                vertical='center',
                wrap_text=column == 4,
            )
        sheet.cell(row=row_number, column=5).number_format = '#,##0'
    total_row = header_row + len(rows) + 1
    sheet.cell(row=total_row, column=4, value='Tổng doanh thu')
    sheet.cell(row=total_row, column=5, value=int(settlement.revenue))
    sheet.cell(row=total_row + 1, column=4, value='Tỷ lệ giảng viên nhận')
    sheet.cell(row=total_row + 1, column=5, value=float(settlement.teacher_share_rate))
    sheet.cell(row=total_row + 1, column=5).number_format = '0%'
    sheet.cell(row=total_row + 2, column=4, value='Tổng tiền giảng viên nhận')
    sheet.cell(row=total_row + 2, column=5, value=int(settlement.teacher_amount))
    for row_number in range(total_row, total_row + 3):
        label_cell = sheet.cell(row=row_number, column=4)
        value_cell = sheet.cell(row=row_number, column=5)
        label_cell.font = Font(bold=True, color='0F172A')
        value_cell.font = Font(bold=True, color='0F172A')
        label_cell.fill = PatternFill('solid', fgColor='E2E8F0')
        value_cell.fill = PatternFill('solid', fgColor='E2E8F0')
        label_cell.alignment = Alignment(horizontal='right', vertical='center')
        value_cell.alignment = Alignment(horizontal='right', vertical='center')
        value_cell.number_format = '#,##0' if row_number != total_row + 1 else '0%'
    sheet.cell(row=total_row + 2, column=4).fill = PatternFill('solid', fgColor='DCFCE7')
    sheet.cell(row=total_row + 2, column=5).fill = PatternFill('solid', fgColor='DCFCE7')
    for column, width in {'A': 8, 'B': 28, 'C': 18, 'D': 32, 'E': 24}.items():
        sheet.column_dimensions[column].width = width
    sheet.auto_filter.ref = f'A{header_row}:E{header_row + len(rows)}'
    sheet.print_area = f'A1:E{total_row + 2}'
    sheet.print_title_rows = f'1:{header_row}'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = sheet.ORIENTATION_LANDSCAPE
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.4, bottom=0.4, header=0.15, footer=0.15)
    sheet.print_options.horizontalCentered = True

    buffer = io.BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    filename = f'quyet_toan_{settlement.teacher.name}_{settlement.classroom.name}_{settlement.pk}.xlsx'
    response['Content-Disposition'] = f"attachment; filename*=UTF-8''{filename}"
    return response




def classroom_detail(request, pk):
    classroom = get_object_or_404(ClassRoom, pk=pk)
    
    # Get teachers teaching in this classroom
    teachers = classroom.teachers.all().order_by('name')
    
    teacher_reports = []
    for t in teachers:
        # Find students in this classroom enrolled under this teacher
        enrollments = Enrollment.objects.filter(
            student__classroom=classroom,
            teacher=t
        ).select_related('student', 'subject').prefetch_related(
            'student__payments__payment_period'
        ).order_by('student__name')
        
        students_list = []
        for e in enrollments:
            student = e.student
            student.enrolled_subject = e.subject.name
            students_list.append(student)

        attach_payment_period_amounts(students_list, teacher=t)
            
        teacher_reports.append({
            'teacher': t,
            'students': students_list,
            'count': len(students_list)
        })
        
    total_paid_all = Payment.objects.filter(classroom=classroom).aggregate(total=Sum('amount'))['total'] or 0
    
    return render(request, 'lms_manager/classroom_detail.html', {
        'classroom': classroom,
        'total_paid_all': total_paid_all,
        'teacher_reports': teacher_reports,
    })



def student_detail(request, pk):
    student = get_object_or_404(Student, pk=pk)
    classroom = student.classroom
    
    # Get all enrollments for this student
    enrollments = Enrollment.objects.filter(student=student).select_related('subject', 'teacher')
    
    # Get all payment periods that match the student's registered subjects and teachers
    from django.db.models import Q
    q_filter = Q()
    for enrollment in enrollments:
        q_filter |= Q(subject_id=enrollment.subject_id, teacher_id=enrollment.teacher_id)
        
    payment_periods = []
    if enrollments:
        payment_periods = list(
            PaymentPeriod.objects.filter(q_filter)
            .select_related('teacher', 'subject')
            .order_by('name', 'subject__name')
        )
    
    # Build report for this student
    records = []
    total_paid = 0
    
    for period in payment_periods:
        # Find payment
        payment = Payment.objects.filter(
            student=student,
            classroom=classroom,
            subject=period.subject,
            payment_period=period
        ).first()
        
        if payment:
            status = "Đã đóng"
            amount = payment.amount
            date = payment.payment_date
            payment_id = payment.id
            total_paid += amount
        else:
            status = "Chưa đóng"
            amount = 0
            date = None
            payment_id = None
            
        records.append({
            'period': period,
            'status': status,
            'amount': amount,
            'date': date,
            'payment_id': payment_id,
            'payment': payment,
        })
        
    return render(request, 'lms_manager/student_detail.html', {
        'student': student,
        'classroom': classroom,
        'enrollments': enrollments,
        'records': records,
        'total_paid': total_paid,
    })


# -------------------------------------------------------------
# PAYMENT PERIOD CRUD
# -------------------------------------------------------------
def payment_period_list(request):
    periods = PaymentPeriod.objects.all().order_by('name')
    return render(request, 'lms_manager/payment_period_list.html', {'periods': periods})

def payment_period_create(request):
    if request.method == 'POST':
        form = PaymentPeriodForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Đã thêm đợt đóng tiền thành công!")
            return redirect('payment_period_list')
    else:
        initial_data = {}
        if 'teacher' in request.GET:
            initial_data['teacher'] = request.GET.get('teacher')
        if 'subject' in request.GET:
            initial_data['subject'] = request.GET.get('subject')
        form = PaymentPeriodForm(initial=initial_data)
    return render(request, 'lms_manager/form.html', {'form': form, 'title': 'Thêm đợt đóng tiền'})

def payment_period_update(request, pk):
    period = get_object_or_404(PaymentPeriod, pk=pk)
    if request.method == 'POST':
        form = PaymentPeriodForm(request.POST, instance=period)
        if form.is_valid():
            form.save()
            messages.success(request, "Đã cập nhật đợt đóng tiền thành công!")
            return redirect('payment_period_list')
    else:
        form = PaymentPeriodForm(instance=period)
    return render(request, 'lms_manager/form.html', {'form': form, 'title': 'Sửa đợt đóng tiền'})

def payment_period_delete(request, pk):
    period = get_object_or_404(PaymentPeriod, pk=pk)
    period.delete()
    messages.success(request, "Đã xóa đợt đóng tiền!")
    return redirect('payment_period_list')


# -------------------------------------------------------------
# NEW VIEWS AND APIS FOR EXCEL IMPORT AND STATUS TOGGLE
# -------------------------------------------------------------
def _parse_import_date(date_str):
    if not date_str:
        return None
    from datetime import datetime
    date_str = date_str.strip()
    for fmt in ('%Y-%m-%d %H:%M:%S', '%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%m/%d/%Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(date_str, fmt).date()
        except ValueError:
            continue
    try:
        from dateutil import parser
        return parser.parse(date_str, dayfirst=True).date()
    except (ImportError, ValueError, TypeError):
        return None


def classroom_import_excel(request, pk):
    classroom = get_object_or_404(ClassRoom, pk=pk)
    teacher_id = request.GET.get('teacher') or request.POST.get('teacher_id')
    teacher = None
    if teacher_id:
        teacher = get_object_or_404(Teacher, pk=teacher_id)
    else:
        teacher = classroom.teachers.order_by('name').first()
    
    if request.method == 'POST':
        if teacher:
            assign_teacher_to_classroom(teacher, classroom)
        excel_file = request.FILES.get('excel_file')
        if not excel_file:
            messages.error(request, "Vui lòng chọn một tệp Excel hoặc CSV để tải lên!")
            return redirect(request.path + (f"?teacher={teacher_id}" if teacher_id else ""))
            
        file_name = excel_file.name.lower()
        rows = []
        try:
            if file_name.endswith('.xlsx') or file_name.endswith('.xls'):
                import openpyxl
                wb = openpyxl.load_workbook(excel_file, read_only=True, data_only=True)
                sheet = wb.active
                for r in sheet.iter_rows(values_only=True):
                    if any(x is not None for x in r):
                        rows.append([str(x).strip() if x is not None else "" for x in r])
            elif file_name.endswith('.csv'):
                import csv
                import io
                file_data = excel_file.read().decode('utf-8-sig')
                io_string = io.StringIO(file_data)
                reader = csv.reader(io_string)
                for r in reader:
                    if r:
                        rows.append([x.strip() for x in r])
            else:
                messages.error(request, "Định dạng tệp không hợp lệ! Vui lòng tải lên tệp .xlsx hoặc .csv.")
                return redirect(request.path + (f"?teacher={teacher_id}" if teacher_id else ""))
                
            if not rows:
                messages.error(request, "Tệp trống hoặc không chứa dữ liệu!")
                return redirect(request.path + (f"?teacher={teacher_id}" if teacher_id else ""))
                
            # Locate the header row.  Downloadable templates may include a
            # title and short instructions above the column headers.
            required = ['Tên', 'SĐT', 'Lớp']
            header_row_index = None
            header_map = None
            headers = []
            for row_index, candidate_row in enumerate(rows[:10]):
                candidate_headers = [cell.strip() for cell in candidate_row]
                candidate_map = {}
                for req in required:
                    for column_index, header in enumerate(candidate_headers):
                        if header.lower() == req.lower():
                            candidate_map[req] = column_index
                            break
                if len(candidate_map) == len(required):
                    header_row_index = row_index
                    header_map = candidate_map
                    headers = candidate_headers
                    break

            if header_row_index is None:
                messages.error(request, f"Tệp tải lên thiếu các cột bắt buộc: {', '.join(required)}")
                return redirect(request.path + (f"?teacher={teacher_id}" if teacher_id else ""))
                
            # Find optional start_date column
            start_date_idx = -1
            for idx, h in enumerate(headers):
                if h.strip().lower() in ['ngày bắt đầu học', 'ngày bắt đầu', 'ngay bat dau hoc', 'ngay bat dau', 'ngày nhập học', 'ngay nhap hoc', 'ngày nhập', 'ngay nhap']:
                    start_date_idx = idx
                    break

            # Process rows
            imported_count = 0
            for row_idx, row in enumerate(rows[header_row_index + 1:], start=header_row_index + 2):
                if not row or len(row) <= max(header_map.values()):
                    continue
                    
                name = row[header_map['Tên']].strip()
                phone = row[header_map['SĐT']].strip()
                class_name = row[header_map['Lớp']].strip()
                
                if not name or not class_name:
                    messages.warning(request, f"Dòng {row_idx}: Bị bỏ quan vì thiếu Tên hoặc Lớp.")
                    continue
                    
                capitalized_name = " ".join(word.capitalize() for word in name.split())
                # This import is launched from a specific classroom.  Keep all
                # imported students in that classroom even if the spreadsheet
                # contains a different class name.
                row_classroom = classroom
                if teacher:
                    assign_teacher_to_classroom(teacher, row_classroom)
                
                start_date_val = None
                if start_date_idx != -1 and len(row) > start_date_idx:
                    start_date_val = _parse_import_date(row[start_date_idx])

                student = Student.objects.create(
                    name=capitalized_name,
                    phone=phone if phone else None,
                    classroom=row_classroom,
                    start_date=start_date_val
                )
                
                if teacher:
                    subject = teacher.subjects.first() or Subject.objects.first()
                    if subject:
                        Enrollment.objects.get_or_create(
                            student=student,
                            subject=subject,
                            defaults={'teacher': teacher}
                        )
                
                imported_count += 1
                
            messages.success(request, f"Đã nhập thành công {imported_count} học sinh vào hệ thống!")
            if teacher_id:
                return redirect('teacher_detail', pk=teacher.id)
            return redirect('classroom_detail', pk=classroom.id)
            
        except Exception as e:
            messages.error(request, f"Đã xảy ra lỗi khi đọc tệp: {str(e)}")
            return redirect(request.path + (f"?teacher={teacher_id}" if teacher_id else ""))
            
    return render(request, 'lms_manager/classroom_import_excel.html', {
        'classroom': classroom,
        'teacher': teacher
    })


from django.views.decorators.http import require_POST

@require_POST
def student_toggle_period(request, pk, period_num):
    # Kept temporarily so old clients receive a clear error instead of a 404.
    # A payment status may only change through the payment-recording workflow.
    return JsonResponse({
        'success': False,
        'error': 'Trạng thái học phí chỉ được cập nhật khi ghi nhận khoản thu.'
    }, status=405)


def get_student_teachers(request, student_id):
    student = get_object_or_404(Student, pk=student_id)
    enrollments = Enrollment.objects.filter(student=student).select_related('teacher')
    teachers = []
    seen = set()
    for e in enrollments:
        if e.teacher.id not in seen:
            seen.add(e.teacher.id)
            teachers.append({
                'id': e.teacher.id,
                'name': e.teacher.name
            })
    return JsonResponse({'teachers': teachers})


def get_student_unpaid_periods(request, student_id):
    student = get_object_or_404(Student, pk=student_id)
    unpaid = []
    for i in range(1, 9):
        status = getattr(student, f"dot_{i}")
        if status == 'Chưa đóng':
            unpaid.append({
                'id': i,
                'name': f"Đợt {i}"
            })
    return JsonResponse({'periods': unpaid})


def get_receipt_url(request, student_id, period_num):
    from django.urls import reverse
    payment = Payment.objects.filter(
        student_id=student_id,
        payment_period__name=f"Đợt {period_num}"
    ).first()
    
    if payment:
        import os
        if not payment.receipt_pdf or not os.path.exists(payment.receipt_pdf.path):
            payment.generate_receipt_pdf()
            payment.refresh_from_db()
            
        return JsonResponse({
            'success': True,
            'receipt_url': reverse('payment_receipt', kwargs={'pk': payment.id}),
            'payment_id': payment.id
        })
    else:
        return JsonResponse({
            'success': False,
            'error': 'Không tìm thấy biên lai đóng tiền cho đợt này.'
        })
