import json
import io
from pathlib import Path
from unittest.mock import patch
from django.test import TestCase, override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.contrib.messages import get_messages
from django.urls import reverse
from django.utils import translation
from lms_manager.models import ClassRoom, Subject, Teacher, Student, Enrollment, PaymentPeriod, Payment, PaymentBatch, TeacherSettlement
from lms_manager.views import assign_teacher_to_classroom


@override_settings(STATICFILES_STORAGE='django.contrib.staticfiles.storage.StaticFilesStorage')
class LMSManagerQueryTest(TestCase):
    def setUp(self):
        self.classroom = ClassRoom.objects.create(name="10A1")
        self.subject = Subject.objects.create(name="Math")
        self.teacher = Teacher.objects.create(name="Mr. Smith")
        self.teacher.subjects.add(self.subject)
        self.teacher.classes.add(self.classroom)
        
        self.student1 = Student.objects.create(name="John Doe", classroom=self.classroom)
        self.student2 = Student.objects.create(name="Jane Doe", classroom=self.classroom)
        
        self.enrollment1 = Enrollment.objects.create(
            student=self.student1,
            subject=self.subject,
            teacher=self.teacher
        )
        self.enrollment2 = Enrollment.objects.create(
            student=self.student2,
            subject=self.subject,
            teacher=self.teacher
        )
        
        self.payment_period = PaymentPeriod.objects.create(
            name="Period 1",
            teacher=self.teacher,
            subject=self.subject
        )
        
        # student1 paid, student2 didn't
        self.payment = Payment.objects.create(
            student=self.student1,
            classroom=self.classroom,
            subject=self.subject,
            teacher=self.teacher,
            payment_period=self.payment_period,
            amount=100000
        )
        
    def test_unpaid_enrollments_query(self):
        period = self.payment_period
        unpaid = Enrollment.objects.filter(
            student__classroom=self.classroom,
            subject=period.subject,
            teacher=period.teacher,
        ).exclude(
            student__payments__payment_period=period
        )
        
        self.assertEqual(unpaid.count(), 1)
        self.assertEqual(unpaid.first().student, self.student2)

    def test_teacher_list_searches_by_name_and_phone(self):
        other_teacher = Teacher.objects.create(name='Ms. Jones', phone='0988111222')

        with translation.override('en'):
            by_name = self.client.get(reverse('teacher_list'), {'q': 'smith'})
            by_phone = self.client.get(reverse('teacher_list'), {'q': '111222'})

        self.assertContains(by_name, self.teacher.name)
        self.assertNotContains(by_name, other_teacher.name)
        self.assertContains(by_phone, other_teacher.name)
        self.assertNotContains(by_phone, self.teacher.name)

    def test_student_and_classroom_exports_are_formatted_xlsx_for_a4_printing(self):
        export_specs = [
            ('student_list_export', []),
            ('student_export', [self.student1.id]),
            ('classroom_list_export', []),
            ('classroom_export', [self.classroom.id]),
        ]
        from openpyxl import load_workbook

        for url_name, args in export_specs:
            with self.subTest(url_name=url_name):
                with translation.override('en'):
                    url = reverse(url_name, args=args)
                    response = self.client.get(url)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(
                    response['Content-Type'],
                    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                )
                worksheet = load_workbook(io.BytesIO(response.content), data_only=True).active
                self.assertEqual(int(worksheet.page_setup.paperSize), int(worksheet.PAPERSIZE_A4))
                self.assertEqual(worksheet.page_setup.fitToWidth, 1)
                self.assertEqual(worksheet.page_setup.fitToHeight, 1)

    def test_student_list_paid_amount_is_derived_from_recorded_payments(self):
        self.payment_period.name = "Đợt 1"
        self.payment_period.save()
        self.student1.dot_1 = "Chưa đóng"
        self.student1.save()
        self.student2.dot_1 = "Đã đóng"
        self.student2.save()

        with translation.override('en'):
            response = self.client.get(reverse('student_list'))

        self.assertContains(response, '100,000 VNĐ', count=1)
        self.assertNotContains(response, '>Đã đóng<')
        self.assertNotContains(response, 'togglePaymentStatus')
        self.assertNotContains(response, 'togglePaymentStatus')

    def test_manual_payment_status_endpoint_cannot_change_payment_data(self):
        with translation.override('en'):
            response = self.client.post(reverse('student_toggle_period', args=[self.student2.id, 1]))

        self.assertEqual(response.status_code, 405)
        self.assertEqual(Payment.objects.filter(student=self.student2).count(), 0)

    def test_batch_payment_page_renders(self):
        with translation.override('en'):
            response = self.client.get(reverse('payment_add'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Thu học phí')

    def test_excel_import_uses_the_selected_classroom_and_teacher(self):
        upload = SimpleUploadedFile(
            'students.csv',
            'Tên,SĐT,Lớp\nNguyen Van C,0900000000,Lớp khác\n'.encode('utf-8'),
            content_type='text/csv',
        )
        with translation.override('en'):
            response = self.client.post(
                f"{reverse('classroom_import_excel', args=[self.classroom.id])}?teacher={self.teacher.id}",
                {'excel_file': upload},
            )

        self.assertEqual(response.status_code, 302)
        student = Student.objects.get(name='Nguyen Van C')
        self.assertEqual(student.classroom, self.classroom)
        self.assertTrue(Enrollment.objects.filter(student=student, teacher=self.teacher).exists())

    def test_downloaded_student_excel_template_can_be_imported(self):
        template_path = Path(__file__).resolve().parents[1] / 'static/templates/Mau_import_hoc_sinh.xlsx'
        upload = SimpleUploadedFile(
            template_path.name,
            template_path.read_bytes(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        with translation.override('en'):
            response = self.client.post(
                reverse('classroom_import_excel', args=[self.classroom.id]),
                {'excel_file': upload},
            )

        self.assertEqual(response.status_code, 302)
        self.assertTrue(
            Student.objects.filter(name='Nguyễn Văn A', classroom=self.classroom).exists(),
            f"students={list(Student.objects.values_list('name', flat=True))}; messages={[str(message) for message in get_messages(response.wsgi_request)]}",
        )

    def test_classroom_is_limited_to_one_teacher_and_enrolls_its_students(self):
        other_subject = Subject.objects.create(name='English')
        other_teacher = Teacher.objects.create(name='Ms. Jones')
        other_teacher.subjects.add(other_subject)
        other_teacher.classes.add(self.classroom)
        unassigned_student = Student.objects.create(name='Sam Doe', classroom=self.classroom)
        Enrollment.objects.create(student=unassigned_student, subject=other_subject, teacher=other_teacher)

        assign_teacher_to_classroom(self.teacher, self.classroom)

        self.assertEqual(list(self.classroom.teachers.all()), [self.teacher])
        self.assertFalse(Enrollment.objects.filter(student=unassigned_student, teacher=other_teacher).exists())
        self.assertTrue(Enrollment.objects.filter(student=unassigned_student, teacher=self.teacher).exists())

    def test_teacher_can_remove_student_from_its_class_without_deleting_student_or_payments(self):
        with translation.override('en'):
            response = self.client.post(
                reverse(
                    'teacher_classroom_student_remove',
                    args=[self.teacher.id, self.classroom.id, self.student1.id],
                ),
            )

        self.assertEqual(response.status_code, 302)
        self.student1.refresh_from_db()
        self.assertIsNone(self.student1.classroom)
        self.assertFalse(Enrollment.objects.filter(student=self.student1, teacher=self.teacher).exists())
        self.assertTrue(Payment.objects.filter(pk=self.payment.pk).exists())

    def test_classroom_detail_shows_paid_amount_instead_of_paid_status(self):
        with translation.override('en'):
            response = self.client.get(reverse('classroom_detail', args=[self.classroom.id]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '100,000 VNĐ')
        self.assertContains(response, 'Đợt 1: đã đóng')
        self.assertContains(response, 'data-paid-1="1"')
        self.assertContains(response, 'data-paid-2="0"')
        self.assertContains(response, 'getAttribute(`data-paid-${period}`)')
        self.assertContains(
            response,
            f'/teachers/{self.teacher.id}/classes/{self.classroom.id}/students/{self.student1.id}/remove/',
        )

    @patch('lms_manager.models.Payment.generate_receipt_pdf')
    def test_teacher_can_settle_multiple_periods_once_and_download_excel(self, _payment_pdf):
        second_period = PaymentPeriod.objects.create(
            name='Period 2', teacher=self.teacher, subject=self.subject
        )
        Payment.objects.create(
            student=self.student2,
            classroom=self.classroom,
            subject=self.subject,
            teacher=self.teacher,
            payment_period=second_period,
            amount=150000,
        )

        with translation.override('en'):
            preview = self.client.get(
                reverse('teacher_class_settlement', args=[self.teacher.id, self.classroom.id]),
                {'period': [1, 2]},
            )
            response = self.client.post(
                reverse('teacher_class_settlement', args=[self.teacher.id, self.classroom.id]),
                {'period': [1, 2]},
            )

        self.assertEqual(preview.status_code, 200)
        self.assertContains(preview, '250,000')
        self.assertContains(preview, 'Đợt 1')
        self.assertContains(preview, 'Đợt 8')
        self.assertContains(preview, '100,000 VNĐ')
        self.assertContains(preview, '150,000 VNĐ')
        self.assertEqual(response.status_code, 302)
        settlement = TeacherSettlement.objects.get()
        self.assertEqual(settlement.revenue, 250000)
        self.assertEqual(settlement.teacher_amount, 200000)
        self.assertEqual(settlement.payments.count(), 2)

        export = self.client.get(response['Location'])
        self.assertEqual(export.status_code, 200)
        self.assertEqual(
            export['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertTrue(export.content.startswith(b'PK'))
        from openpyxl import load_workbook
        worksheet = load_workbook(io.BytesIO(export.content), data_only=True).active
        self.assertEqual(worksheet['A1'].value, 'QUYẾT TOÁN GIẢNG VIÊN - Mr. Smith')
        self.assertEqual(worksheet['E8'].value, 250000)
        self.assertEqual(worksheet['E10'].value, 200000)
        self.assertEqual(int(worksheet.page_setup.paperSize), int(worksheet.PAPERSIZE_A4))
        self.assertEqual(worksheet.page_setup.orientation, worksheet.ORIENTATION_LANDSCAPE)
        self.assertEqual(worksheet.page_setup.fitToWidth, 1)
        self.assertEqual(worksheet.page_setup.fitToHeight, 1)

        with translation.override('en'):
            after_settlement = self.client.get(
                reverse('teacher_class_settlement', args=[self.teacher.id, self.classroom.id]),
            )
        self.assertContains(after_settlement, 'Đợt 1')
        self.assertContains(after_settlement, 'Đợt 8')

        with translation.override('en'):
            settled_preview = self.client.get(
                reverse('teacher_class_settlement', args=[self.teacher.id, self.classroom.id]),
                {'period': [1]},
            )
        self.assertContains(settled_preview, 'Đã quyết toán')
        self.assertNotContains(settled_preview, 'Xác nhận quyết toán &amp; tải Excel', html=True)

    @patch('lms_manager.models.Payment.generate_receipt_pdf')
    def test_late_payment_in_a_settled_period_can_be_settled_without_old_payments(self, _payment_pdf):
        with translation.override('en'):
            self.client.post(
                reverse('teacher_class_settlement', args=[self.teacher.id, self.classroom.id]),
                {'period': [1]},
            )
        first_settlement = TeacherSettlement.objects.get()
        self.assertEqual(first_settlement.payments.count(), 1)

        late_payment = Payment.objects.create(
            student=self.student2,
            classroom=self.classroom,
            subject=self.subject,
            teacher=self.teacher,
            payment_period=self.payment_period,
            amount=120000,
        )
        with translation.override('en'):
            preview = self.client.get(
                reverse('teacher_class_settlement', args=[self.teacher.id, self.classroom.id]),
                {'period': [1]},
            )
            self.client.post(
                reverse('teacher_class_settlement', args=[self.teacher.id, self.classroom.id]),
                {'period': [1]},
            )

        self.assertContains(preview, '120,000')
        self.assertEqual(TeacherSettlement.objects.count(), 2)
        late_settlement = TeacherSettlement.objects.exclude(pk=first_settlement.pk).get()
        self.assertEqual(list(late_settlement.payments.all()), [late_payment])
        self.assertEqual(late_settlement.revenue, 120000)

    def test_settlement_shows_empty_message_for_a_period_without_payments(self):
        with translation.override('en'):
            response = self.client.get(
                reverse('teacher_class_settlement', args=[self.teacher.id, self.classroom.id]),
                {'period': [8]},
            )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Đợt 1')
        self.assertContains(response, 'Đợt 8')
        self.assertContains(response, 'Không có khoản thu chưa quyết toán trong các đợt đã chọn.')

    @patch('lms_manager.models.PaymentBatch.generate_receipt_pdf')
    @patch('lms_manager.models.Payment.generate_receipt_pdf')
    def test_batch_payment_creates_individual_payments_and_one_batch_receipt(self, _payment_pdf, _batch_pdf):
        with translation.override('en'):
            response = self.client.post(
                reverse('payment_add'),
                {
                    'payment_date': '2026-07-19',
                    'batch_items': json.dumps([
                        {
                            'student_id': self.student2.id,
                            'teacher_id': self.teacher.id,
                            'period_nums': [1, 2],
                            'amount': 500000,
                        },
                    ]),
                },
                HTTP_X_REQUESTED_WITH='XMLHttpRequest',
            )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['success'])
        batch = PaymentBatch.objects.get()
        self.assertEqual(batch.payments.count(), 2)
        self.assertEqual(
            set(batch.payments.values_list('payment_period__name', flat=True)),
            {'Đợt 1', 'Đợt 2'},
        )
        self.assertTrue(all(payment.student == self.student2 for payment in batch.payments.all()))
        self.assertTrue(all(payment.amount == 500000 for payment in batch.payments.all()))
